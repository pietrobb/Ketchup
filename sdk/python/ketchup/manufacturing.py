"""Opt-in manufacturing exports; adapters are explicitly supplied TRUSTED Python objects.

No plugin discovery/import by name occurs. The neutral manifest requires only the
standard library. JafWebCutAdapter alone lazily requires openpyxl>=3.1,<4.
Adapters must render without external side effects; this is not a Python sandbox.
"""
from __future__ import annotations

from collections.abc import Iterable, Mapping
from copy import deepcopy
import ctypes
import errno
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import stat
import sys
import tempfile
from typing import Protocol


class ManufacturingError(ValueError):
    """Invalid production job, adapter configuration or unsafe output."""


class ManufacturingAdapter(Protocol):
    id: str

    def render(self, job: dict) -> Mapping[str, bytes]: ...


def _text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ManufacturingError(f"{label} must be a nonempty string")
    return value


def _json(value):
    try:
        return json.dumps(value, ensure_ascii=True, allow_nan=False, sort_keys=True,
                          indent=2).encode("utf-8") + b"\n"
    except (TypeError, ValueError, RecursionError) as exc:
        raise ManufacturingError("job must contain finite JSON data") from exc


def _validate_job(job):
    if not isinstance(job, dict) or job.get("schema") != "ketchup.production-job.v1":
        raise ManufacturingError("expected ketchup.production-job.v1 job")
    for key in ("document_id", "source_revision"):
        if type(job.get(key)) is not int or job[key] < 0:
            raise ManufacturingError(f"{key} must be a nonnegative integer")
    _text(job.get("source_digest"), "source_digest")
    parts = job.get("parts")
    if not isinstance(parts, list) or not parts:
        raise ManufacturingError("parts must be a nonempty list")
    codes, paths = set(), set()
    for part in parts:
        if not isinstance(part, dict):
            raise ManufacturingError("part must be an object")
        code = part.get("code")
        if not isinstance(code, str) or not re.fullmatch(r"[A-Z0-9_-]{1,64}", code):
            raise ManufacturingError("part code must be 1..64 uppercase A-Z0-9_- characters")
        if code in codes:
            raise ManufacturingError(f"duplicate part code: {code}")
        codes.add(code)
        if not isinstance(part.get("instance_path"), dict):
            raise ManufacturingError("instance_path must be an object")
        path = _json(part["instance_path"])
        if path in paths:
            raise ManufacturingError("duplicate physical instance_path")
        paths.add(path)
        _text(part.get("name"), "part name")
        _text(part.get("material_key"), "material_key")
        dims = part.get("dimensions_mm")
        if (not isinstance(dims, list) or len(dims) != 3
                or any(type(d) not in (int, float) or not math.isfinite(d) or d <= 0
                       for d in dims)):
            raise ManufacturingError("dimensions_mm must contain positive finite length, width, thickness")
    if not isinstance(job.get("outputs"), dict):
        raise ManufacturingError("outputs must be an object keyed by adapter id")
    for adapter_id in job["outputs"]:
        _text(adapter_id, "output adapter id")
    _json(job)


def _filename(name):
    # Flat portable namespace: no directories, ADS, device names or normalization aliases.
    if (not isinstance(name, str) or not re.fullmatch(r"[A-Za-z0-9_-][A-Za-z0-9_.-]{0,199}", name)
            or name.endswith(".") or name.split(".")[0].upper() in
            {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(10)),
             *(f"LPT{i}" for i in range(10))}):
        raise ManufacturingError(f"unsafe filename: {name!r}")
    return name


def _check_destination(destination):
    path = Path(destination).absolute()
    _filename(path.name)
    for parent in (path.parent, *path.parent.parents):
        info = parent.lstat()  # Missing parents are deliberately not created.
        if (stat.S_ISLNK(info.st_mode)
                or getattr(info, "st_file_attributes", 0) & 0x400):
            raise ManufacturingError("symlink/junction destination parents are forbidden")
        if not stat.S_ISDIR(info.st_mode):
            raise ManufacturingError("destination parent must be a directory")
    if os.path.lexists(path):
        raise FileExistsError(f"destination already exists: {path}")
    return path


def _publish(stage, destination):
    """Atomic no-replace directory rename (fail closed on unsupported platforms)."""
    if os.name == "nt":
        os.rename(stage, destination)  # Windows rename never replaces an existing target.
    elif sys.platform.startswith("linux"):
        libc = ctypes.CDLL(None, use_errno=True)
        rename = getattr(libc, "renameat2", None)
        if rename is None:
            raise ManufacturingError("atomic no-replace publication unavailable")
        rename.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p,
                           ctypes.c_uint]
        rename.restype = ctypes.c_int
        if rename(-100, os.fsencode(stage), -100, os.fsencode(destination), 1):
            error = ctypes.get_errno()
            raise OSError(error, os.strerror(error), str(destination))
    else:
        raise ManufacturingError("atomic no-replace publication unsupported on this platform")


def export_job(job: dict, destination, adapters: Iterable[ManufacturingAdapter], *,
               confirmed: bool = False) -> dict:
    """Publish a NEW directory and return its audit manifest (also manifest.json).

    confirmed must be literally True. Existing parents must be real directories,
    owned/trusted by the caller (not concurrently replaced by another process).
    All renders, validations and file writes precede an atomic no-replace rename.
    An empty adapters list exports only the neutral manifest. Unselected outputs
    are recorded in the source job but never written as machine files.
    """
    if confirmed is not True:
        raise ManufacturingError("manufacturing export requires explicit confirmed=True")
    _validate_job(job)
    snapshot = deepcopy(job)
    target = _check_destination(destination)
    selected = list(adapters)
    ids = set()
    for adapter in selected:
        ident = _text(getattr(adapter, "id", None), "adapter id")
        if ident in ids or not callable(getattr(adapter, "render", None)):
            raise ManufacturingError("duplicate adapter id or missing render(job)")
        ids.add(ident)
    files, names, records = {}, {"manifest.json"}, []
    for adapter in selected:
        rendered = adapter.render(deepcopy(snapshot))
        if not isinstance(rendered, Mapping):
            raise ManufacturingError("adapter render must return a filename-to-bytes mapping")
        for name, content in rendered.items():
            _filename(name)
            if name.casefold() in names:
                raise ManufacturingError(f"filename collision: {name}")
            if not isinstance(content, bytes):
                raise ManufacturingError(f"adapter output must be bytes: {name}")
            names.add(name.casefold())
            files[name] = content
            records.append({"adapter_id": adapter.id, "filename": name,
                            "size_bytes": len(content),
                            "sha256": hashlib.sha256(content).hexdigest()})
    manifest = {"schema": "ketchup.manufacturing-export.v1", "status": "completed",
                "job": snapshot, "job_sha256": hashlib.sha256(_json(snapshot)).hexdigest(),
                "adapters": [a.id for a in selected], "files": records}
    files["manifest.json"] = _json(manifest)
    stage = Path(tempfile.mkdtemp(prefix=".ketchup-export-", dir=target.parent))
    try:
        for name, content in files.items():
            with (stage / name).open("xb") as stream:
                stream.write(content)
                stream.flush()
                os.fsync(stream.fileno())
        _check_destination(target)
        _publish(stage, target)
    finally:
        if stage.exists():
            shutil.rmtree(stage)
    return manifest


class HomagWoodwopAdapter:
    """Pass through explicitly requested ASCII MPRs, without rewriting any byte.

    Validates the transport envelope, identity and ASCII content, not machine
    toolpaths or physical machine safety. Outputs may be a subset of the parts.
    """
    id = "homag-woodwop4"

    def render(self, job: dict) -> Mapping[str, bytes]:
        _validate_job(job)
        parts = {p["code"]: p for p in job["parts"]}
        if any(not re.fullmatch(r"[A-Z0-9_-]{12}", code) for code in parts):
            raise ManufacturingError("HOMAG requires exact12 character part codes")
        outputs = job["outputs"].get(self.id)
        if not isinstance(outputs, list):
            raise ManufacturingError("explicit homag-woodwop4 outputs list is required")
        result, seen = {}, set()
        for output in outputs:
            if not isinstance(output, dict):
                raise ManufacturingError("HOMAG output must be an object")
            code = output.get("code")
            if not isinstance(code, str) or code not in parts or code in seen:
                raise ManufacturingError("unknown or duplicate HOMAG output code")
            seen.add(code)
            name = output.get("filename")
            if name != code + ".mpr":
                raise ManufacturingError("HOMAG filename must equal part code + .mpr")
            _filename(name)
            content = _text(output.get("content"), "MPR content")
            if any(ord(ch) > 127 or (ord(ch) < 32 and ch not in "\r\n\t")
                   or ord(ch) == 127 for ch in content):
                raise ManufacturingError("MPR content must be ASCII text without control characters")
            result[name] = content.encode("ascii")
        return result


class JafWebCutAdapter:
    """FurniGen's original D..U template mapping, one physical part per row.

    material_mapping is material_key -> exact JAF material string (no guessing).
    allow_rotation is an explicit bool: True leaves L blank (Export H=1), False
    writes x (Export H=0). Dimensions retain their order and fractional mm; this
    template has no integer-only dimension validation. Codes are limited to 40
    characters by the template. Export formulas are materialized as values, so
    this is a finalized order snapshot, not an automatically recalculating form.
    """
    id = "jaf-webcut"
    filename = "JAF_WebCut_v50.xlsx"

    def __init__(self, material_mapping: Mapping[str, str], *, allow_rotation: bool):
        if not isinstance(material_mapping, Mapping) or type(allow_rotation) is not bool:
            raise ManufacturingError("explicit material mapping and boolean allow_rotation required")
        self.material_mapping = dict(material_mapping)
        for key, value in self.material_mapping.items():
            _text(key, "material key")
            _text(value, "JAF material")
        self.allow_rotation = allow_rotation

    def render(self, job: dict) -> Mapping[str, bytes]:
        _validate_job(job)
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("JafWebCutAdapter requires optional openpyxl>=3.1,<4") from exc
        workbook = load_workbook(Path(__file__).parent / "templates" / self.filename)
        try:
            sheet, export = workbook["všeobecný"], workbook["Export"]
            capacity = 0
            for row in range(18, sheet.max_row + 1):
                if (sheet.cell(row, 3).data_type != "f"
                        or export.cell(row - 16, 15).data_type != "f"):
                    break
                capacity += 1
            if len(job["parts"]) > capacity:
                raise ManufacturingError(f"JAF template row capacity exceeded ({capacity})")
            for row in sheet.iter_rows(min_row=18, min_col=4, max_col=21):
                for cell in row:
                    cell.value = None
            # Remove stale formula results and phantom unused rows from the export grid.
            for row in export.iter_rows(min_row=2, max_col=16):
                for cell in row:
                    cell.value = None
            group, position, last_group = 0, 0, None
            for offset, part in enumerate(job["parts"]):
                code = part["code"]
                if len(code) > 40:
                    raise ManufacturingError("JAF template barcode accepts at most 40 characters")
                material = self.material_mapping.get(part["material_key"])
                if material is None:
                    raise ManufacturingError(f"missing JAF material mapping: {part['material_key']}")
                length, width, thickness = part["dimensions_mm"]
                key = material, thickness
                if key != last_group:
                    group, position, last_group = group + 1, 0, key
                position += 1
                r = 18 + offset
                values = {4: group, 5: str(position), 6: material, 7: thickness,
                          8: part["name"], 9: 1, 10: length, 11: width,
                          12: None if self.allow_rotation else "x", 19: code, 20: code}
                for column, value in values.items():
                    self._cell(sheet.cell(r, column), value)
                # Exact equivalents of the original Export A..P formulas for these inputs.
                combined = material + " " + format(thickness, ".15g")
                out = [group, str(position), combined, part["name"], 1, length, width,
                       int(self.allow_rotation), None, None, None, None, None, code, code, None]
                for column, value in enumerate(out, 1):
                    self._cell(export.cell(offset + 2, column), value)
            buffer = io.BytesIO()
            workbook.save(buffer)
            return {self.filename: buffer.getvalue()}
        finally:
            workbook.close()

    @staticmethod
    def _cell(cell, value):
        if isinstance(value, str):
            if len(value) > 32767 or any(ord(c) < 32 and c not in "\t\r\n" for c in value):
                raise ManufacturingError("JAF cell text contains illegal characters or is too long")
            cell.value = value
            cell.data_type = "s"  # Never interpret untrusted names/materials as formulas.
            cell.number_format = "@"
        else:
            cell.value = value
