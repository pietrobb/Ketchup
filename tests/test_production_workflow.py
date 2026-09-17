"""Real saved-document -> optional production adapters, without running a CNC."""
import io
import os
from pathlib import Path
import sys

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "sdk" / "python"))
from ketchup import Session, HeadlessError
from ketchup.manufacturing import HomagWoodwopAdapter, JafWebCutAdapter


def feature(result, kind):
    return next(f["id"] for f in result["state"]["features"]
                if f["id"] in result["created"]["feature_ids"] and f["kind"] == kind)


def test_saved_codes_drive_excel_and_mpr_via_optional_plugins(tmp_path):
    executable = os.environ.get("KETCHUP_HEADLESS")
    if not executable:
        pytest.skip("Set KETCHUP_HEADLESS to run the real production workflow")
    from openpyxl import load_workbook
    with Session(executable=executable) as session:
        doc = session.new_document()
        first = doc.box("Bok", 400, 200, 18)
        definition, = first["created"]["definition_ids"]
        first_id, = first["created"]["occurrence_ids"]
        base = feature(first, "Pad")
        circle = doc.create_sketch(definition, "Dowel drill", [
            {"type": "circle", "id": 1, "center_mm": [50, 25], "radius_mm": 4},
        ], workplane={"type": "frame", "origin_mm": [0, 200, 18],
                      "x_axis": [1, 0, 0], "y_axis": [0, -1, 0]})
        doc.pocket(definition, "Blind drilling", base, feature(circle, "Sketch"), 8)
        second = doc.box("Polica", 400, 200, 18, translation_mm=(600, 0, 0))
        second_id, = second["created"]["occurrence_ids"]
        doc.apply([
            {"operation": "upsert_classification_dimension", "dimension_id": 900000,
             "name": "ketchup.fabrication-role.v1", "categories": [
                 {"id": 900001, "name": "fabrication.timber-member.v1"}]},
            {"operation": "set_occurrence_classification", "dimension_id": 900000,
             "category_id": 900001,
             "selector": {"type": "occurrences", "occurrence_ids": [first_id, second_id]}},
        ])
        assignments = [
            {"instance_path": {"root_occurrence_id": ident, "steps": []}, "code": code}
            for ident, code in [(first_id, "000000004201"), (second_id, "000000004202")]
        ]
        doc.set_production_codes(assignments)
        assert doc.production_codes() == assignments
        with pytest.raises(HeadlessError) as missing_save:
            doc.production_job()
        assert missing_save.value.code == "unsaved_changes"
        path = tmp_path / "furniture.ketchup"
        doc.save(path)
        doc = session.open_document(path)
        assert doc.production_codes() == assignments
        neutral = doc.production_job()
        assert neutral["outputs"] == {}
        assert len(neutral["parts"]) == 2
        material = neutral["parts"][0]["material_key"]
        adapters = [JafWebCutAdapter({material: "TEST BOARD"}, allow_rotation=False),
                    HomagWoodwopAdapter()]
        with pytest.raises(ValueError, match="confirmed"):
            doc.export_production(tmp_path / "order", adapters,
                                  machine_adapters=["homag-woodwop4"])
        result = doc.export_production(tmp_path / "order", adapters,
                                      machine_adapters=["homag-woodwop4"], confirmed=True)
        assert result["job"]["source_digest"] == doc.state["canonical_digest"]
        programs = list((tmp_path / "order").glob("*.mpr"))
        assert [p.name for p in programs] == ["000000004201.mpr"]
        assert b"\\BohrVert\\" in programs[0].read_bytes()
        workbook = load_workbook(io.BytesIO((tmp_path / "order" / "JAF_WebCut_v50.xlsx").read_bytes()), data_only=True)
        try:
            for row, assignment in enumerate(assignments, 18):
                assert workbook["všeobecný"].cell(row, 19).value == assignment["code"]
                assert workbook["Export"].cell(row - 16, 14).value == assignment["code"]
                assert workbook["všeobecný"].cell(row, 20).value == (assignment["code"] if row == 18 else None)
                assert workbook["Export"].cell(row - 16, 15).value == (assignment["code"] if row == 18 else None)
                assert workbook["všeobecný"].cell(row, 9).value == 1
        finally:
            workbook.close()
        # The same saved job can target another workflow without any HOMAG output.
        class OtherMachine:
            id = "other-machine"
            def render(self, job):
                assert job["outputs"] == {}
                return {"other-process.txt": "\n".join(p["code"] for p in job["parts"]).encode("ascii")}
        doc.export_production(tmp_path / "other-order", [OtherMachine()], confirmed=True)
        assert (tmp_path / "other-order" / "other-process.txt").read_text().splitlines() == [a["code"] for a in assignments]
        with pytest.raises(FileExistsError):
            doc.export_production(tmp_path / "order", [], confirmed=True)
