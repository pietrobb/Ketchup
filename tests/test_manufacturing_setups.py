"""Setup-code transport is not machine orientation/toolpath qualification."""
from copy import deepcopy
import io

import pytest

from test_manufacturing_plugins import job
from ketchup.manufacturing import (
    export_job, HomagWoodwopAdapter, JafWebCutAdapter, ManufacturingError,
    assign_setup_codes,
)


def two_setups():
    source = job(1)
    part = source["parts"][0]
    part["operations"].append({"kind": "circular-drill", "operation_id": "drill-B"})
    part["machining_setups"].append({
        "id": "B", "code": "000000000002", "operation_ids": ["drill-B"],
        "dowel_hole_ids": [],
    })
    source["outputs"]["homag-woodwop4"].append({
        "part_code": part["code"], "setup_id": "B", "code": "000000000002",
        "filename": "000000000002.mpr", "content": "SECOND SETUP TEST FIXTURE\r\n",
    })
    return source


def test_two_setup_codes_stay_on_one_physical_row_and_match_programs(tmp_path):
    from openpyxl import load_workbook
    source = two_setups()
    original = deepcopy(source)
    assigned = assign_setup_codes(source, {source["parts"][0]["code"]: {
        "A": "00107001", "B": "00107002"}})
    assert source == original
    jaf = JafWebCutAdapter({"oak": "BOARD"}, allow_rotation=False,
                          second_barcode_confirmed=True)
    manifest = export_job(assigned, tmp_path / "order", [
        jaf, HomagWoodwopAdapter(program_code_length=8)], confirmed=True)
    assert {p.name for p in (tmp_path / "order").glob("*.mpr")} == {"00107001.mpr", "00107002.mpr"}
    assert (tmp_path / "order" / "00107002.mpr").read_bytes() == b"SECOND SETUP TEST FIXTURE\r\n"
    assert manifest["job"]["parts"][0]["code"] == source["parts"][0]["code"]
    for data_only in [True, False]:
        wb = load_workbook(io.BytesIO((tmp_path / "order" / jaf.filename).read_bytes()), data_only=data_only)
        try:
            ws, ex = wb["všeobecný"], wb["Export"]
            assert ws["I18"].value == ex["E2"].value == 1
            assert ws["S18"].value == ex["N2"].value == source["parts"][0]["code"]
            for left, right, code in [("T18", "O2", "00107001"), ("U18", "P2", "00107002")]:
                assert ws[left].value == ex[right].value == code
                assert ws[left].data_type == ex[right].data_type == "s"
            assert ws["I19"].value is None and ex["E3"].value is None
        finally:
            wb.close()


def test_second_barcode_needs_explicit_jaf_confirmation(tmp_path):
    with pytest.raises(ManufacturingError, match="second barcode"):
        export_job(two_setups(), tmp_path / "order", [
            JafWebCutAdapter({"oak": "BOARD"}, allow_rotation=False)], confirmed=True)
    assert not list(tmp_path.iterdir())


def test_cut_only_has_physical_id_but_no_program_barcode():
    from openpyxl import load_workbook
    adapter = JafWebCutAdapter({"oak": "BOARD"}, allow_rotation=False)
    wb = load_workbook(io.BytesIO(adapter.render(job())[adapter.filename]), data_only=True)
    try:
        assert wb["všeobecný"]["S19"].value == "000000000001"
        assert wb["Export"]["N3"].value == "000000000001"
        assert all(wb[s][c].value is None for s, c in [
            ("všeobecný", "T19"), ("všeobecný", "U19"), ("Export", "O3"), ("Export", "P3")])
    finally:
        wb.close()


@pytest.mark.parametrize("change", [
    lambda j: j["parts"][0]["machining_setups"][1].update(code=j["parts"][0]["machining_setups"][0]["code"]),
    lambda j: j["parts"][0]["machining_setups"][1].update(id="A"),
    lambda j: j["parts"][0]["machining_setups"][1].update(operation_ids=[]),
    lambda j: j["parts"][0]["machining_setups"][1].update(operation_ids=["UNKNOWN"]),
    lambda j: j["parts"][0]["machining_setups"][1].update(operation_ids=["drill-A", "drill-B"]),
    lambda j: j["parts"][0]["machining_setups"][1].update(dowel_hole_ids=["UNKNOWN"]),
    lambda j: j["parts"][0].pop("machining_setups"),
    lambda j: j["parts"][0]["machining_setups"][1].update(code="../escape"),
    lambda j: j["parts"][0]["machining_setups"][1].update(code="=FORMULA"),
])
def test_invalid_setup_partition_never_publishes(tmp_path, change):
    source = two_setups()
    change(source)
    with pytest.raises(ManufacturingError):
        export_job(source, tmp_path / "order", [], confirmed=True)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("change", [
    lambda j: j["outputs"]["homag-woodwop4"].pop(),
    lambda j: j["outputs"]["homag-woodwop4"][1].update(setup_id="A"),
    lambda j: j["outputs"]["homag-woodwop4"][1].update(part_code="WRONG"),
    lambda j: j["outputs"]["homag-woodwop4"][1].update(code="000000000003"),
    lambda j: j["outputs"]["homag-woodwop4"][1].update(filename="other.mpr"),
])
def test_missing_or_mismatched_second_program_never_publishes(tmp_path, change):
    source = two_setups()
    change(source)
    with pytest.raises(ManufacturingError):
        export_job(source, tmp_path / "order", [HomagWoodwopAdapter()], confirmed=True)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("assignments", [
    {"UNKNOWN": {"A": "000000000003"}},
    {"000000000000": {"B": "000000000003"}},
    {"000000000000": {"A": "../evil"}},
    {"000000000000": {"A": "SHORT", "B": "SHORT"}},
    {"000000000001": {"A": "000000000003"}},
    {"000000000000": []},
    [],
])
def test_code_assignment_cannot_create_setups_or_change_source(assignments):
    source = job()
    original = deepcopy(source)
    with pytest.raises(ManufacturingError):
        assign_setup_codes(source, assignments)
    assert source == original


def test_neutral_setup_ids_generalize_beyond_two_faces():
    source = two_setups()
    source["parts"][0]["machining_setups"][1]["id"] = "fixture-3"
    source["outputs"] = {}
    # Machine-neutral data allows other setup conventions; JAF does not silently truncate them.
    assign_setup_codes(source, {source["parts"][0]["code"]: {"fixture-3": "CUSTOM_3"}})
    with pytest.raises(ManufacturingError, match="A/B"):
        JafWebCutAdapter({"oak": "BOARD"}, allow_rotation=False,
                         second_barcode_confirmed=True).render(source)


@pytest.mark.parametrize("length", [True, 0, 65, "8", None])
def test_invalid_program_code_length(length):
    with pytest.raises(ManufacturingError):
        HomagWoodwopAdapter(program_code_length=length)


def test_short_program_name_requires_explicit_destination_length():
    source = job(1)
    source = assign_setup_codes(source, {source["parts"][0]["code"]: {"A": "00107001"}})
    with pytest.raises(ManufacturingError, match="exactly 12"):
        HomagWoodwopAdapter().render(source)
    assert list(HomagWoodwopAdapter(program_code_length=8).render(source)) == ["00107001.mpr"]


def test_second_setup_dowel_coverage_cannot_be_silently_omitted():
    source = two_setups()
    part = source["parts"][0]
    part["operations"].pop()
    part["dowel_holes"] = [{"kind": "dowel_drill", "hole_id": "joint-B-hole"}]
    part["machining_setups"][1].update(operation_ids=[], dowel_hole_ids=["joint-B-hole"])
    assert len(HomagWoodwopAdapter().render(source)) == 2
    part["machining_setups"].pop()
    with pytest.raises(ManufacturingError, match="cover every"):
        HomagWoodwopAdapter().render(source)


def test_setup_program_codes_are_globally_unique_not_just_per_part(tmp_path):
    source = job()
    other = source["parts"][1]
    other["operations"].append({"kind": "circular-drill", "operation_id": "other-drill"})
    other["machining_setups"] = [{"id": "A", "code": "000000000000",
                                  "operation_ids": ["other-drill"], "dowel_hole_ids": []}]
    with pytest.raises(ManufacturingError, match="duplicate.*program code"):
        export_job(source, tmp_path / "order", [], confirmed=True)
    assert not list(tmp_path.iterdir())
