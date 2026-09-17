"""Native setup-A export and explicit refusal of unqualified JAF operations."""
import json
import os

import pytest

from test_production_workflow import feature
from ketchup import Session, HeadlessError
from ketchup.manufacturing import HomagWoodwopAdapter, JafWebCutAdapter, ManufacturingError


def make_document(session, tmp_path, machining="blind"):
    doc = session.new_document()
    first = doc.box("Panel", 400, 200, 18)
    definition, = first["created"]["definition_ids"]
    occurrence, = first["created"]["occurrence_ids"]
    base = feature(first, "Pad")
    # These are geometry fixtures, NOT a confirmed machine flip convention.
    frame = {"type": "frame", "origin_mm": [0, 200, 18],
             "x_axis": [1, 0, 0], "y_axis": [0, -1, 0]}
    if machining == "bottom":
        frame = {"type": "frame", "origin_mm": [0, 0, 0],
                 "x_axis": [1, 0, 0], "y_axis": [0, 1, 0]}
    entities = [{"type": "circle", "id": 1, "center_mm": [50, 25], "radius_mm": 4}]
    if machining == "through-pocket":
        points = [[20, 20], [120, 20], [120, 100], [20, 100]]
        entities = [{"type": "line", "id": i + 1, "start_mm": p,
                     "end_mm": points[(i + 1) % 4]} for i, p in enumerate(points)]
    sketch = doc.create_sketch(definition, "Machining", entities, workplane=frame)
    doc.pocket(definition, "Machining cut", base, feature(sketch, "Sketch"),
               18 if machining.startswith("through") else 8)
    doc.apply([
        {"operation": "upsert_classification_dimension", "dimension_id": 900000,
         "name": "ketchup.fabrication-role.v1", "categories": [
             {"id": 900001, "name": "fabrication.timber-member.v1"}]},
        {"operation": "set_occurrence_classification", "dimension_id": 900000,
         "category_id": 900001,
         "selector": {"type": "occurrences", "occurrence_ids": [occurrence]}},
    ])
    doc.set_production_codes([{"instance_path": {"root_occurrence_id": occurrence, "steps": []},
                               "code": "PANEL-42"}])
    path = tmp_path / "panel.ketchup"
    doc.save(path)
    return session.open_document(path)


def test_native_physical_identity_is_independent_of_program_code_and_saved_manifest(tmp_path):
    executable = os.environ.get("KETCHUP_HEADLESS")
    if not executable:
        pytest.skip("Set KETCHUP_HEADLESS for native setup regressions")
    from openpyxl import load_workbook
    with Session(executable=executable) as session:
        doc = make_document(session, tmp_path)
        original_digest = doc.state["canonical_digest"]
        neutral = doc.production_job()
        assert neutral["schema"] == "ketchup.production-job.v2"
        part, = neutral["parts"]
        assert part["code"] == "PANEL-42"
        setup, = part["machining_setups"]
        assert setup == {"id": "A", "code": "PANEL-42",
                         "operation_ids": [part["operations"][1]["operation_id"]],
                         "dowel_hole_ids": []}
        # Missing setup cannot be invented merely by assigning a code.
        with pytest.raises(ManufacturingError, match="unknown setup"):
            doc.export_production(tmp_path / "invented", [], confirmed=True,
                                  setup_codes={"PANEL-42": {"B": "00107002"}})
        assert not (tmp_path / "invented").exists()
        before = doc.production_job(machine_adapters=["homag-woodwop4"])
        manifest = doc.export_production(tmp_path / "order", [
            JafWebCutAdapter({part["material_key"]: "TEST BOARD"}, allow_rotation=False),
            HomagWoodwopAdapter(program_code_length=8),
        ], machine_adapters=["homag-woodwop4"], confirmed=True,
            setup_codes={"PANEL-42": {"A": "00107001"}})
        assert doc.state["canonical_digest"] == original_digest
        assert doc.production_codes()[0]["code"] == "PANEL-42"
        assert manifest["job"]["parts"][0]["machining_setups"][0]["code"] == "00107001"
        assert json.loads((tmp_path / "order" / "manifest.json").read_bytes()) == manifest
        assert not list((tmp_path / "order").glob("*.svg"))
        assert [p.name for p in (tmp_path / "order").glob("*.mpr")] == ["00107001.mpr"]
        assert (tmp_path / "order" / "00107001.mpr").read_bytes() == before["outputs"]["homag-woodwop4"][0]["content"].encode("ascii")
        wb = load_workbook(tmp_path / "order" / "JAF_WebCut_v50.xlsx", data_only=True)
        try:
            assert [wb["všeobecný"].cell(18, c).value for c in [19, 20, 21]] == ["PANEL-42", "00107001", None]
            assert [wb["Export"].cell(2, c).value for c in [14, 15, 16]] == ["PANEL-42", "00107001", None]
        finally:
            wb.close()


@pytest.mark.parametrize("machining", ["bottom", "through-drill", "through-pocket"])
def test_native_unqualified_jaf_machining_never_publishes(tmp_path, machining):
    executable = os.environ.get("KETCHUP_HEADLESS")
    if not executable:
        pytest.skip("Set KETCHUP_HEADLESS for native setup regressions")
    with Session(executable=executable) as session:
        doc = make_document(session, tmp_path, machining)
        neutral = doc.production_job()
        assert len(neutral["parts"][0]["operations"]) == 2
        with pytest.raises(HeadlessError) as rejected:
            doc.export_production(tmp_path / "unsafe", [HomagWoodwopAdapter()],
                                  machine_adapters=["homag-woodwop4"],
                                  vertical_pocket_tool_number=101 if machining == "through-pocket" else None,
                                  confirmed=True, setup_codes={"PANEL-42": {"A": "000000000042"}})
        assert rejected.value.code == "production_blocked"
        assert not (tmp_path / "unsafe").exists()
