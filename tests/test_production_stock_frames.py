"""Neutral stock frames and adapter-specific cutting constraints; no physical machines."""
import io
import json
import os

import pytest

from test_manufacturing_plugins import job
from ketchup import Session, HeadlessError
from ketchup.manufacturing import export_job, JafWebCutAdapter, ManufacturingError


@pytest.mark.parametrize("shape", [None, "profile_extrusion", "unknown", True])
def test_jaf_requires_explicit_rectangular_stock(tmp_path, shape):
    source = job()
    source["parts"][0]["stock_shape"] = shape
    with pytest.raises(ManufacturingError, match="rectangular"):
        export_job(source, tmp_path / "order",
                   [JafWebCutAdapter({"oak": "TEST BOARD"}, allow_rotation=False)],
                   confirmed=True)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("polygon", [False, True])
@pytest.mark.parametrize("frame", [
    {"type": "frame", "origin_mm": [10, 20, 30],
     "x_axis": [1, 0, 0], "y_axis": [0, 1, 0]},
    {"type": "frame", "origin_mm": [10, 20, 30],
     "x_axis": [0, 1, 0], "y_axis": [0, 0, 1]},
])
def test_native_neutral_preserves_stock_frame_and_profile(tmp_path, polygon, frame):
    executable = os.environ.get("KETCHUP_HEADLESS")
    if not executable:
        pytest.skip("Set KETCHUP_HEADLESS for native stock-frame regression")
    from openpyxl import load_workbook
    points = [[5, 7], [405, 7], [305 if polygon else 405, 207], [5, 207]]
    entities = [{"type": "line", "id": i + 1, "start_mm": start,
                 "end_mm": points[(i + 1) % len(points)]} for i, start in enumerate(points)]
    with Session(executable=executable) as session:
        doc = session.new_document()
        created = doc.apply([{"operation": "create_part", "name": "Framed stock",
                              "workplane": frame, "entities": entities, "constraints": [],
                              "feature": {"type": "extrusion", "distance_mm": 18},
                              "translation_mm": [0, 0, 0]}])
        occurrence, = created["created"]["occurrence_ids"]
        doc.apply([
            {"operation": "upsert_classification_dimension", "dimension_id": 900000,
             "name": "ketchup.fabrication-role.v1", "categories": [
                 {"id": 900001, "name": "fabrication.timber-member.v1"}]},
            {"operation": "set_occurrence_classification", "dimension_id": 900000,
             "category_id": 900001,
             "selector": {"type": "occurrences", "occurrence_ids": [occurrence]}},
        ])
        doc.set_production_codes([{"instance_path": {"root_occurrence_id": occurrence, "steps": []},
                                   "code": "000000000001"}])
        path = tmp_path / "stock.ketchup"
        doc.save(path)
        doc = session.open_document(path)
        neutral = doc.production_job()
        part, = neutral["parts"]
        assert part["dimensions_mm"] == [400, 200, 18], part
        assert part["operations"][0]["geometry"]["frame"]["origin_mm"] == frame["origin_mm"], part
        assert part["stock_shape"] == ("profile_extrusion" if polygon else "rectangular_prism")
        assert part["coordinate_frame"] == "definition_local_mm"
        x, y = frame["x_axis"], frame["y_axis"]
        z = [x[1]*y[2]-x[2]*y[1], x[2]*y[0]-x[0]*y[2], x[0]*y[1]-x[1]*y[0]]
        assert part["stock_frame"] == {
            "origin_mm": [frame["origin_mm"][i] + 5*x[i] + 7*y[i] for i in range(3)],
            "axes": [x, y, z],
        }
        geometry = part["operations"][0]["geometry"]
        assert geometry["frame"]["origin_mm"] == frame["origin_mm"]
        assert geometry["frame"]["x_axis"] == x
        assert geometry["frame"]["y_axis"] == y
        assert geometry["start_mm"] == frame["origin_mm"]
        assert geometry["length_axis"] == z
        assert geometry["cross_section"] == [
            {"kind": "line", "start_mm": p, "end_mm": points[(i+1) % 4]}
            for i, p in enumerate(points)]

        class OtherMachine:
            id = "profile-aware-machine"

            def render(self, source):
                assert source["parts"][0] == part
                assert source["outputs"] == {}
                return {"stock.json": json.dumps(source["parts"][0]).encode()}

        result = doc.export_production(tmp_path / "other", [OtherMachine()], confirmed=True)
        assert json.loads((tmp_path / "other" / "stock.json").read_bytes()) == part
        assert result["job"]["parts"] == [part]
        jaf = JafWebCutAdapter({part["material_key"]: "TEST BOARD"}, allow_rotation=False)
        if polygon:
            with pytest.raises(ManufacturingError, match="rectangular"):
                doc.export_production(tmp_path / "jaf", [jaf], confirmed=True)
            assert not (tmp_path / "jaf").exists()
        else:
            doc.export_production(tmp_path / "jaf", [jaf], confirmed=True)
            workbook = load_workbook(io.BytesIO((tmp_path / "jaf" / jaf.filename).read_bytes()), data_only=True)
            try:
                assert [workbook["všeobecný"].cell(18, c).value for c in [10, 11, 7, 19, 20]] == [400, 200, 18, part["code"], None]
            finally:
                workbook.close()
        with pytest.raises(HeadlessError) as rejected:
            doc.production_job(machine_adapters=["homag-woodwop4"])
        assert rejected.value.code == "production_blocked"
